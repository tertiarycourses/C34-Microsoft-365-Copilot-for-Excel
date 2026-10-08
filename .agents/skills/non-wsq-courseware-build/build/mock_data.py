"""Deterministic synthetic HR / recruitment datasets for the C34 labs.

Scenario: BrightPath Talent Solutions — a FICTIONAL Singapore recruitment,
flexible-staffing and HR-outsourcing firm. Every company, person, ID and value
in these workbooks is synthetic and generated from fixed random seeds, so the
expected answers quoted in each lab's "Test it" step are reproducible.

`build_lab_data(num, folder)` writes the lab's workbook(s) into <folder>/data/.
`EXPECTED[num]` holds the values the lab text quotes for verification.
"""
import os
import random
import statistics
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COMPANY = "BrightPath Talent Solutions"
FICTION_NOTE = (
    "All organisations, people, IDs and values in this workbook are FICTIONAL training data "
    "created for the Microsoft 365 Copilot for Excel course (C34). Any resemblance to a real "
    "company or person is coincidental."
)
REPORT_DATE = date(2026, 6, 30)

SECTORS = ["Banking & Finance", "Healthcare", "Logistics", "Technology", "Manufacturing", "Retail & F&B"]
ROLES = {
    "Banking & Finance": [("Relationship Manager", "Permanent", 6500), ("Compliance Analyst", "Permanent", 5800),
                          ("Accounts Executive", "Contract", 3800)],
    "Healthcare": [("Staff Nurse", "Permanent", 4200), ("Patient Service Associate", "Contract", 2600),
                   ("Clinic Assistant", "Temp", 2300)],
    "Logistics": [("Warehouse Supervisor", "Permanent", 3600), ("Logistics Coordinator", "Contract", 3100),
                  ("Forklift Operator", "Temp", 2400)],
    "Technology": [("Software Engineer", "Permanent", 7200), ("IT Support Engineer", "Contract", 4000),
                   ("Data Analyst", "Permanent", 5600)],
    "Manufacturing": [("Production Technician", "Contract", 2900), ("Quality Engineer", "Permanent", 4800),
                      ("Maintenance Technician", "Permanent", 3400)],
    "Retail & F&B": [("Retail Associate", "Temp", 2100), ("Store Manager", "Permanent", 4000),
                     ("Customer Service Executive", "Contract", 2700)],
}
SOURCES = ["LinkedIn", "JobStreet", "MyCareersFuture", "Referral", "BrightPath Website", "Job Fair", "Talent Database"]
SOURCE_W = [25, 20, 15, 10, 12, 6, 12]
SOURCE_QUALITY = {"LinkedIn": 1.0, "JobStreet": 0.85, "MyCareersFuture": 0.8, "Referral": 1.6,
                  "BrightPath Website": 1.05, "Job Fair": 0.7, "Talent Database": 1.25}
SOURCE_COST = {"LinkedIn": 38, "JobStreet": 22, "MyCareersFuture": 6, "Referral": 55,
               "BrightPath Website": 9, "Job Fair": 31, "Talent Database": 12}
RECRUITERS = [
    ("R01", "Aisha Rahman", "Finance & Tech Desk", 6),
    ("R02", "Daniel Tan", "Finance & Tech Desk", 5),
    ("R03", "Priya Nair", "Healthcare Desk", 5),
    ("R04", "Marcus Lim", "Industrial Desk", 6),
    ("R05", "Siti Hamid", "Industrial Desk", 5),
    ("R06", "Kevin Ong", "Retail Desk", 4),
]
RECRUITER_SECTORS = {
    "R01": ["Banking & Finance", "Technology"], "R02": ["Technology", "Banking & Finance"],
    "R03": ["Healthcare"], "R04": ["Logistics", "Manufacturing"], "R05": ["Manufacturing", "Logistics"],
    "R06": ["Retail & F&B"],
}
CLIENTS = {
    "Banking & Finance": ["Kestrel Wealth", "Orchard Crest Bank", "Silverline Insurance"],
    "Healthcare": ["Azurepoint Health", "Lumora Clinics", "Meridian Care Group"],
    "Logistics": ["Tanjong Freight", "Corallis Logistics", "Bayfront Cold Chain"],
    "Technology": ["Quillon Systems", "Nimbusly Cloud", "Arcwave Analytics"],
    "Manufacturing": ["Pinnacrest Manufacturing", "Veritek Precision", "Solano Polymers"],
    "Retail & F&B": ["Halcyon Retail", "Verdane Foods", "Kopi Lane Cafes"],
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")


def _recruiter_for(rng, sector):
    ids = [r for r, secs in RECRUITER_SECTORS.items() if sector in secs]
    return rng.choice(ids)


def _rand_date(rng, start, end):
    return start + timedelta(days=rng.randint(0, (end - start).days))


# --------------------------------------------------------------------------- workbook helpers
def _write_rows(ws, headers, rows, table_name=None, formats=None, widths=None, start_row=1):
    formats = formats or {}
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        if table_name is None:
            cell.font = Font(bold=True)
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            fmt = formats.get(headers[c - 1])
            if fmt:
                cell.number_format = fmt
            elif isinstance(v, date):
                cell.number_format = "yyyy-mm-dd"
    last = f"{get_column_letter(len(headers))}{start_row + len(rows)}"
    if table_name:
        t = Table(displayName=table_name, ref=f"A{start_row}:{last}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    for c, h in enumerate(headers, 1):
        w = (widths or {}).get(h) or max(11, min(34, len(h) + 3))
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _about_sheet(wb, title, purpose, dictionary):
    ws = wb.create_sheet("About_This_Data", 0)
    ws["A1"] = f"{COMPANY} — {title}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = FICTION_NOTE
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:C2")
    ws.row_dimensions[2].height = 48
    ws["A3"] = purpose
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:C3")
    ws.row_dimensions[3].height = 48
    ws["A5"], ws["B5"], ws["C5"] = "Sheet / Column", "Type", "Meaning"
    for c in ("A5", "B5", "C5"):
        ws[c].fill = HEADER_FILL
        ws[c].font = HEADER_FONT
    for i, (col, typ, meaning) in enumerate(dictionary, 6):
        ws.cell(row=i, column=1, value=col)
        ws.cell(row=i, column=2, value=typ)
        ws.cell(row=i, column=3, value=meaning).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90


def _prompt_log_sheet(wb):
    ws = wb.create_sheet("Prompt_Log")
    headers = ["Step", "Mode (Chat / Plan / Edit)", "Prompt", "What Copilot did", "How I verified", "Decision (Keep / Revise / Undo)"]
    _write_rows(ws, headers, [], None, widths={"Prompt": 60, "What Copilot did": 40, "How I verified": 40})
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).fill = HEADER_FILL
        ws.cell(row=1, column=c).font = HEADER_FONT


def _save(wb, folder, name):
    out = os.path.join(folder, "data")
    os.makedirs(out, exist_ok=True)
    if "Sheet" in wb.sheetnames and wb["Sheet"].max_row == 1 and wb["Sheet"]["A1"].value is None:
        del wb["Sheet"]
    path = os.path.join(out, name)
    wb.save(path)
    return path


# =========================================================================== LAB 1
def _lab1():
    rng = random.Random(3401)
    rows = []
    start, end = date(2026, 1, 5), date(2026, 6, 26)
    for i in range(1, 241):
        sector = rng.choices(SECTORS, [16, 15, 17, 20, 14, 18])[0]
        title, etype, base = rng.choice(ROLES[sector])
        source = rng.choices(SOURCES, SOURCE_W)[0]
        rec = _recruiter_for(rng, sector)
        applied = _rand_date(rng, start, end)
        yrs = max(0, round(rng.gauss(5, 3)))
        exp_sal = int(round(base * (0.85 + yrs * 0.03) * rng.uniform(0.92, 1.12), -1))
        q = SOURCE_QUALITY[source]
        age = (REPORT_DATE - applied).days
        r = rng.random()
        if age < 10 and r < 0.6:
            stage = "Applied"
        elif rng.random() < 0.42 / q:
            stage = "Rejected"
        elif rng.random() < 0.07:
            stage = "Withdrawn"
        else:
            s = rng.random() * q
            stage = "Screening" if s < 0.25 else "Interview" if s < 0.62 else "Offer" if s < 0.80 else "Hired"
        score = None
        if stage in ("Interview", "Offer", "Hired"):
            score = {"Interview": rng.choice([2, 3, 3, 4]), "Offer": rng.choice([3, 4, 4, 5]),
                     "Hired": rng.choice([4, 4, 5, 5])}[stage]
        elif stage == "Rejected" and rng.random() < 0.3:
            score = rng.choice([1, 2, 2, 3])
        days = min(age, max(1, int(rng.gauss({"Applied": 4, "Screening": 9, "Interview": 18, "Offer": 26,
                                                 "Hired": 33, "Rejected": 12, "Withdrawn": 15}[stage], 5))))
        rows.append([f"APP-{i:04d}", applied, f"JOB-{1000 + ROLES[sector].index((title, etype, base)) + 10 * SECTORS.index(sector)}",
                     title, sector, etype, source, dict((r[0], r[1]) for r in RECRUITERS)[rec], yrs, exp_sal,
                     stage, score, days])
    rows.sort(key=lambda r: (r[1], r[0]))
    # renumber in date order so IDs read naturally
    for i, r in enumerate(rows, 1):
        r[0] = f"APP-{i:04d}"
    # deliberate data-quality issues for the lab to find (no row is deleted)
    variant_rows = [r for r in rows if r[6] == "LinkedIn"][3:6]
    for r in variant_rows:
        r[6] = "Linked In"
    js_rows = [r for r in rows if r[6] == "JobStreet"][4:6]
    for r in js_rows:
        r[6] = "JobStreet "
    dup_src = [r for r in rows if r[10] == "Rejected"][7:9]
    dups = [list(r) for r in dup_src]
    rows = rows + dups
    headers = ["Application_ID", "Applied_Date", "Job_ID", "Job_Title", "Client_Sector", "Employment_Type", "Source",
               "Recruiter", "Years_Experience", "Expected_Salary", "Stage", "Interview_Score", "Days_in_Pipeline"]

    def clean(s):
        return s.strip().replace("Linked In", "LinkedIn")
    uniq = {}
    for r in rows:
        uniq.setdefault(r[0], r)
    ur = list(uniq.values())
    by_source = {}
    for r in ur:
        by_source[clean(r[6])] = by_source.get(clean(r[6]), 0) + 1
    hired = [r for r in ur if r[10] == "Hired"]
    hires_by_source = {}
    for r in hired:
        hires_by_source[clean(r[6])] = hires_by_source.get(clean(r[6]), 0) + 1
    by_sector = {}
    for r in ur:
        by_sector[r[4]] = by_sector.get(r[4], 0) + 1
    tech_interview = sum(1 for r in ur if r[4] == "Technology" and r[10] == "Interview")
    exp = dict(
        rows=len(rows), unique=len(ur), dup_ids=sorted({r[0] for r in dups}),
        hired=len(hired), linkedin_clean=by_source["LinkedIn"], linkedin_raw_exact=sum(1 for r in ur if r[6] == "LinkedIn"),
        variants_linkedin=len(variant_rows), variants_jobstreet=len(js_rows),
        top_sector=max(by_sector, key=by_sector.get), top_sector_n=max(by_sector.values()),
        top_hire_source=max(hires_by_source, key=hires_by_source.get), top_hire_source_n=max(hires_by_source.values()),
        tech_interview=tech_interview, first_date=min(r[1] for r in rows), last_date=max(r[1] for r in rows),
        score5=sum(1 for r in ur if r[11] == 5),
        linkedin_clean_all=sum(1 for r in rows if clean(r[6]) == "LinkedIn"),
        jobstreet_clean_all=sum(1 for r in rows if clean(r[6]) == "JobStreet"),
        tech_interview_all=sum(1 for r in rows if r[4] == "Technology" and r[10] == "Interview"),
        hired_all=sum(1 for r in rows if r[10] == "Hired"),
        dup_detail=[(r[0], r[6], r[10]) for r in dups],
    )
    return headers, rows, exp


def build_lab1(folder):
    headers, rows, _ = _lab1()
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    _write_rows(ws, headers, rows, None, {"Expected_Salary": '"S$"#,##0'}, {"Job_Title": 26, "Client_Sector": 18})
    _about_sheet(wb, "Candidate Pipeline (Jan–Jun 2026)",
                 "Purpose: one row per candidate application received by BrightPath's permanent and flexible "
                 "staffing desks. Use it in Lab 1 to explore, question, highlight, sort and filter pipeline data with Copilot. "
                 "The data is deliberately NOT yet formatted as an Excel table, and contains a few quality issues to find.",
                 [("Application_ID", "Text", "Unique application key (APP-0001 …). Should not repeat."),
                  ("Applied_Date", "Date", "Date the application was received."),
                  ("Job_ID / Job_Title", "Text", "Job order the candidate applied for."),
                  ("Client_Sector", "Text", "Sector of the hiring client."),
                  ("Employment_Type", "Text", "Permanent, Contract or Temp."),
                  ("Source", "Text", "Channel the application came from. Canonical values: " + ", ".join(SOURCES) + "."),
                  ("Recruiter", "Text", "BrightPath recruiter who owns the application."),
                  ("Years_Experience", "Number", "Candidate's years of relevant experience."),
                  ("Expected_Salary", "S$ / month", "Candidate's expected monthly base salary."),
                  ("Stage", "Text", "Current stage: Applied, Screening, Interview, Offer, Hired, Rejected, Withdrawn."),
                  ("Interview_Score", "1–5", "Panel score; blank if the candidate was not interviewed."),
                  ("Days_in_Pipeline", "Days", "Days since application to current stage or close.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Candidate_Pipeline.xlsx")]


# =========================================================================== LAB 2
def _lab2():
    rng = random.Random(3402)
    reqs = []
    for i in range(1, 61):
        sector = rng.choices(SECTORS, [17, 15, 16, 20, 15, 17])[0]
        title, _etype, base = rng.choice(ROLES[sector])
        client = rng.choice(CLIENTS[sector])
        rec = _recruiter_for(rng, sector)
        openings = rng.choices([1, 1, 1, 2, 2, 3, 4, 5], k=1)[0]
        open_d = _rand_date(rng, date(2026, 1, 2), date(2026, 6, 15))
        budget = int(round(base * rng.uniform(0.95, 1.25), -2))
        fee = rng.choice([0.18, 0.20, 0.20, 0.22, 0.25])
        r = rng.random()
        age = (REPORT_DATE - open_d).days
        if age > 25 and r < 0.55:
            status = "Filled"
        elif r < 0.62:
            status = "Cancelled"
        elif r < 0.70:
            status = "On Hold"
        else:
            status = "Open"
        if status == "Filled":
            filled = openings
            close = open_d + timedelta(days=max(9, min(age, int(rng.gauss(34 if sector != "Technology" else 46, 10)))))
        elif status == "Cancelled":
            filled = rng.choice([0, 0, min(1, openings - 1)])
            close = open_d + timedelta(days=max(5, min(age, rng.randint(10, 50))))
        else:
            filled = rng.randint(0, max(0, openings - 1))
            close = None
        reqs.append([f"JOB-{2000 + i}", client, sector, title, rec, openings, filled, open_d, close, status, budget, fee])
    reqs.sort(key=lambda r: r[7])
    for i, r in enumerate(reqs, 1):
        r[0] = f"JOB-{2000 + i}"
    headers = ["Job_ID", "Client", "Sector", "Job_Title", "Recruiter_ID", "Openings", "Filled", "Open_Date",
               "Close_Date", "Status", "Monthly_Salary_Budget", "Fee_Rate"]

    def days_open(r):
        return ((r[8] or REPORT_DATE) - r[7]).days
    filled_rows = [r for r in reqs if r[9] == "Filled"]
    fee_total = sum(r[6] * r[10] * 12 * r[11] for r in reqs)
    overdue = [r for r in reqs if r[9] == "Open" and days_open(r) > 45]
    names = {r[0]: r[1] for r in RECRUITERS}
    fee_by_rec = {}
    for r in reqs:
        fee_by_rec[names[r[4]]] = fee_by_rec.get(names[r[4]], 0) + r[6] * r[10] * 12 * r[11]
    sample = next(r for r in reqs if r[9] == "Filled" and r[6] >= 2)
    exp = dict(
        rows=len(reqs), filled_n=len(filled_rows),
        avg_days_filled=round(statistics.mean(days_open(r) for r in filled_rows), 1),
        fee_total=round(fee_total, 2), overdue_n=len(overdue), overdue_ids=[r[0] for r in overdue],
        top_rec=max(fee_by_rec, key=fee_by_rec.get), top_rec_fee=round(max(fee_by_rec.values()), 2),
        openings=sum(r[5] for r in reqs), filled_total=sum(r[6] for r in reqs),
        sample_id=sample[0], sample_fee=round(sample[6] * sample[10] * 12 * sample[11], 2),
        sample_days=days_open(sample), sample_detail=(sample[6], sample[10], sample[11]),
        status_counts={s: sum(1 for r in reqs if r[9] == s) for s in ("Open", "Filled", "On Hold", "Cancelled")},
    )
    return headers, reqs, exp


def build_lab2(folder):
    headers, rows, _ = _lab2()
    wb = Workbook()
    ws = wb.active
    ws.title = "Requisitions"
    _write_rows(ws, headers, rows, "tblRequisitions",
                {"Monthly_Salary_Budget": '"S$"#,##0', "Fee_Rate": "0%"}, {"Client": 24, "Job_Title": 26})
    ws2 = wb.create_sheet("Recruiters")
    _write_rows(ws2, ["Recruiter_ID", "Recruiter_Name", "Desk", "Monthly_Target_Placements"],
                [list(r) for r in RECRUITERS], "tblRecruiters", widths={"Desk": 22})
    ws3 = wb.create_sheet("Control")
    ws3["A1"], ws3["B1"] = "Report_Date", REPORT_DATE
    ws3["B1"].number_format = "yyyy-mm-dd"
    ws3["A1"].font = Font(bold=True)
    ws3["A3"] = "Use this sheet for your independent check formulas (Lab 2, Step 6)."
    ws3.column_dimensions["A"].width = 34
    ws3.column_dimensions["B"].width = 16
    _about_sheet(wb, "Permanent-Recruitment Job Requisitions (H1 2026)",
                 "Purpose: one row per client job order (requisition) on BrightPath's permanent-recruitment desk. BrightPath "
                 "charges a placement fee equal to Fee_Rate x first-year base salary (Monthly_Salary_Budget x 12) for every "
                 "position filled. Use it in Lab 2 to generate and verify formula columns and lookups with Copilot.",
                 [("Job_ID", "Text", "Unique requisition key."), ("Client / Sector", "Text", "Fictional hiring client and its sector."),
                  ("Recruiter_ID", "Text", "Owner; look up the name in tblRecruiters."),
                  ("Openings / Filled", "Count", "Positions requested and positions filled so far."),
                  ("Open_Date / Close_Date", "Date", "Close_Date is blank while the requisition is Open or On Hold."),
                  ("Status", "Text", "Open, Filled, On Hold or Cancelled."),
                  ("Monthly_Salary_Budget", "S$ / month", "Client's budgeted monthly base salary per position."),
                  ("Fee_Rate", "%", "Placement fee as a % of first-year base salary."),
                  ("Control!B1 Report_Date", "Date", "Fixed reporting date 2026-06-30 used for ageing open requisitions.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Job_Requisitions.xlsx")]


# =========================================================================== LAB 3
def _months(y0, m0, n):
    out = []
    y, m = y0, m0
    for _ in range(n):
        out.append(date(y, m, 1))
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def _lab3():
    rng = random.Random(3403)
    months = _months(2025, 1, 18)
    base = {"Banking & Finance": 120, "Healthcare": 150, "Logistics": 170, "Technology": 210,
            "Manufacturing": 130, "Retail & F&B": 190}
    rows = []
    for mi, m in enumerate(months):
        season = 1 + 0.12 * (m.month in (1, 3, 9)) - 0.10 * (m.month in (6, 12))
        for s in SECTORS:
            growth = 1 + (0.025 * mi if s in ("Healthcare", "Technology") else 0.004 * mi)
            if s == "Retail & F&B" and m.year == 2026:
                growth *= 0.82
            apps = int(base[s] * season * growth * rng.uniform(0.9, 1.1))
            screened = int(apps * rng.uniform(0.45, 0.55))
            interviews = int(screened * rng.uniform(0.38, 0.48))
            offers = int(interviews * rng.uniform(0.32, 0.42))
            hire_ratio = rng.uniform(0.70, 0.86) if s != "Technology" else rng.uniform(0.58, 0.70)
            hires = max(1, int(round(offers * hire_ratio)))
            tth = round(rng.gauss({"Technology": 41, "Banking & Finance": 37, "Healthcare": 33}.get(s, 24), 3), 1)
            cph = int(round(rng.gauss({"Technology": 2400, "Banking & Finance": 2100}.get(s, 1100), 120), -1))
            rows.append([m, m.strftime("%Y-%m"), s, apps, screened, interviews, offers, hires, tth, cph])
    headers = ["Month", "Month_Label", "Sector", "Applications", "Screened", "Interviews", "Offers", "Hires",
               "Avg_Time_to_Hire_Days", "Cost_per_Hire"]
    src_rows = []
    rng2 = random.Random(34031)
    for s in SOURCES:
        apps = int(1400 * SOURCE_W[SOURCES.index(s)] / 100 * rng2.uniform(0.9, 1.1))
        hires = int(round(apps * 0.045 * SOURCE_QUALITY[s] * rng2.uniform(0.9, 1.1)))
        spend = int(round(apps * SOURCE_COST[s] * rng2.uniform(0.95, 1.05), -2))
        src_rows.append([s, apps, hires, spend])
    h1 = [r for r in rows if r[0].year == 2026]
    by_sector = {}
    for r in h1:
        by_sector[r[2]] = by_sector.get(r[2], 0) + r[7]
    by_month = {}
    for r in rows:
        by_month[r[1]] = by_month.get(r[1], 0) + r[7]
    funnel = [sum(r[i] for r in h1) for i in (3, 4, 5, 6, 7)]
    hire_rate = {r[0]: r[2] / r[1] for r in src_rows}
    cps = {r[0]: r[3] / r[2] for r in src_rows}
    exp = dict(rows=len(rows), h1_hires=sum(r[7] for r in h1), top_sector=max(by_sector, key=by_sector.get),
               top_sector_hires=max(by_sector.values()), by_sector=by_sector, peak_month=max(by_month, key=by_month.get),
               peak_month_hires=max(by_month.values()), funnel=funnel,
               best_rate_source=max(hire_rate, key=hire_rate.get), best_rate=round(max(hire_rate.values()) * 100, 1),
               lowest_cps_source=min(cps, key=cps.get), lowest_cps=round(min(cps.values())),
               total_2025=sum(r[7] for r in rows if r[0].year == 2025))
    return headers, rows, src_rows, exp


def build_lab3(folder):
    headers, rows, src_rows, _ = _lab3()
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly_Metrics"
    _write_rows(ws, headers, rows, "tblMonthly", {"Month": "mmm-yyyy", "Cost_per_Hire": '"S$"#,##0',
                                                  "Avg_Time_to_Hire_Days": "0.0"}, {"Sector": 18})
    ws2 = wb.create_sheet("Source_Summary")
    _write_rows(ws2, ["Source", "Applications", "Hires", "Sourcing_Spend"], src_rows, "tblSources",
                {"Sourcing_Spend": '"S$"#,##0'}, {"Source": 20})
    wb.create_sheet("Charts")
    _about_sheet(wb, "Monthly Recruitment Metrics (Jan 2025 – Jun 2026)",
                 "Purpose: one row per month per client sector with funnel counts and efficiency measures, plus a "
                 "Jan–Jun 2026 sourcing-channel summary. Use it in Lab 3 to create and refine charts with Copilot. "
                 "Place finished charts on the Charts sheet.",
                 [("Month / Month_Label", "Date / Text", "First day of the month and its yyyy-mm label."),
                  ("Sector", "Text", "Client sector."),
                  ("Applications → Hires", "Count", "Funnel counts for the month: Applications, Screened, Interviews, Offers, Hires."),
                  ("Avg_Time_to_Hire_Days", "Days", "Average days from application to accepted offer for hires that month."),
                  ("Cost_per_Hire", "S$", "Average sourcing + recruiter cost per hire."),
                  ("Source_Summary sheet", "Table", "Applications, hires and sourcing spend by channel, Jan–Jun 2026.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Recruitment_Metrics.xlsx")]


# =========================================================================== LAB 4
LEVELS = ["Junior", "Mid", "Senior", "Lead"]


def _lab4():
    rng = random.Random(3404)
    funcs = {"Banking & Finance": "Finance & Accounting", "Healthcare": "Healthcare Support",
             "Logistics": "Supply Chain", "Technology": "IT & Engineering",
             "Manufacturing": "Production & Quality", "Retail & F&B": "Customer Operations"}
    sector_rate = {"Banking & Finance": 34, "Healthcare": 24, "Logistics": 19, "Technology": 44,
                   "Manufacturing": 21, "Retail & F&B": 16}
    level_mult = {"Junior": 0.75, "Mid": 1.0, "Senior": 1.3, "Lead": 1.62}
    level_exp = {"Junior": (1, 2), "Mid": (4, 2), "Senior": (8, 2.5), "Lead": (13, 3)}
    rows = []
    for i in range(1, 181):
        s = rng.choices(SECTORS, [15, 17, 18, 20, 15, 15])[0]
        lvl = rng.choices(LEVELS, [32, 36, 22, 10])[0]
        mu, sd = level_exp[lvl]
        yrs = max(0, round(rng.gauss(mu, sd), 1))
        pay = round(sector_rate[s] * level_mult[lvl] * (1 + 0.012 * yrs) * rng.uniform(0.88, 1.12), 2)
        markup = rng.uniform(1.22, 1.40) if s != "Technology" else rng.uniform(1.18, 1.30)
        if i in (37, 121):  # two deliberate high-rate outliers for the box plot discussion
            pay = round(pay * 1.9, 2)
        bill = round(pay * markup, 2)
        weeks = rng.choice([12, 13, 24, 26, 26, 39, 52])
        start = _rand_date(rng, date(2025, 7, 1), date(2026, 6, 15))
        ext = "Yes" if rng.random() < (0.35 + 0.1 * (lvl in ("Senior", "Lead"))) else "No"
        rating = round(min(5, max(1, rng.gauss(3.9, 0.6))), 1)
        rows.append([f"PL-{5000 + i}", start, s, funcs[s], lvl, yrs, pay, bill, weeks, ext, rating])
    rows.sort(key=lambda r: r[1])
    for i, r in enumerate(rows, 1):
        r[0] = f"PL-{5000 + i}"
    headers = ["Placement_ID", "Start_Date", "Sector", "Job_Function", "Job_Level", "Years_Experience",
               "Pay_Rate_Hourly", "Bill_Rate_Hourly", "Contract_Weeks", "Extended", "Client_Rating"]
    med = {lv: round(statistics.median(r[7] for r in rows if r[4] == lv), 2) for lv in LEVELS}
    cnt = {lv: sum(1 for r in rows if r[4] == lv) for lv in LEVELS}
    margins = [(r[7] - r[6]) / r[7] for r in rows]
    by_sector_margin = {}
    for s in SECTORS:
        ms = [(r[7] - r[6]) / r[7] for r in rows if r[2] == s]
        by_sector_margin[s] = round(statistics.mean(ms) * 100, 1)
    xs = [r[5] for r in rows]
    ys = [r[6] for r in rows]
    corr = statistics.correlation(xs, ys)
    pays = sorted(r[6] for r in rows)
    q1, _q2, q3 = statistics.quantiles(pays, n=4, method="inclusive")
    upper = q3 + 1.5 * (q3 - q1)
    exp = dict(rows=len(rows), median_bill=med, counts=cnt, avg_margin=round(statistics.mean(margins) * 100, 1),
               sector_margin=by_sector_margin, low_margin_sector=min(by_sector_margin, key=by_sector_margin.get),
               corr=round(corr, 2), pay_outliers=sum(1 for p in pays if p > upper), upper_fence=round(upper, 2),
               max_pay_id=max(rows, key=lambda r: r[6])[0], max_pay=max(r[6] for r in rows))
    return headers, rows, exp


def build_lab4(folder):
    headers, rows, _ = _lab4()
    wb = Workbook()
    ws = wb.active
    ws.title = "Placements"
    _write_rows(ws, headers, rows, "tblPlacements",
                {"Pay_Rate_Hourly": '"S$"#,##0.00', "Bill_Rate_Hourly": '"S$"#,##0.00', "Years_Experience": "0.0",
                 "Client_Rating": "0.0"}, {"Sector": 18, "Job_Function": 22})
    wb.create_sheet("Python_Charts")
    _about_sheet(wb, "Contract Placements (Jul 2025 – Jun 2026)",
                 "Purpose: one row per contract associate placed with a client by BrightPath's flexible-staffing desk. "
                 "BrightPath pays the associate Pay_Rate_Hourly and bills the client Bill_Rate_Hourly. Use it in Lab 4 for "
                 "distribution, relationship and comparison charts with Copilot and Python in Excel.",
                 [("Placement_ID", "Text", "Unique placement key."), ("Start_Date", "Date", "Assignment start date."),
                  ("Sector / Job_Function", "Text", "Client sector and job family."),
                  ("Job_Level", "Text", "Junior, Mid, Senior or Lead."),
                  ("Years_Experience", "Years", "Associate's relevant experience."),
                  ("Pay_Rate_Hourly / Bill_Rate_Hourly", "S$ / hour", "Rate paid to the associate / billed to the client."),
                  ("Gross margin (derived)", "%", "(Bill_Rate_Hourly − Pay_Rate_Hourly) ÷ Bill_Rate_Hourly — not stored; compute it."),
                  ("Contract_Weeks", "Weeks", "Initial contract length."), ("Extended", "Yes/No", "Whether the client extended the contract."),
                  ("Client_Rating", "1–5", "Client's rating of the associate.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Contract_Placements.xlsx")]


# =========================================================================== LAB 5
def _lab5():
    rng = random.Random(3405)
    rows = []
    rec_names = {r[0]: r[1] for r in RECRUITERS}
    for i in range(1, 601):
        s = rng.choices(SECTORS, [16, 15, 17, 20, 14, 18])[0]
        src = rng.choices(SOURCES, SOURCE_W)[0]
        rec = _recruiter_for(rng, s)
        applied = _rand_date(rng, date(2026, 1, 2), date(2026, 6, 20))
        q = SOURCE_QUALITY[src]
        stage = "Applied"
        if rng.random() < 0.55 * min(1.3, q):
            stage = "Screened"
            if rng.random() < 0.45 * min(1.3, q):
                stage = "Interviewed"
                if rng.random() < 0.42:
                    stage = "Offered"
        offer_made = "Yes" if stage == "Offered" else "No"
        accepted, hire_date, tth = "", None, None
        if offer_made == "Yes":
            p_acc = 0.62 if s == "Technology" else 0.80
            if rng.random() < p_acc:
                accepted = "Yes"
                tth = max(12, int(rng.gauss({"Technology": 40, "Banking & Finance": 36, "Healthcare": 31}.get(s, 24), 6)))
                hire_date = applied + timedelta(days=tth)
                if hire_date > REPORT_DATE:
                    accepted, hire_date, tth = "Pending", None, None
            else:
                accepted = "No"
        final = "Hired" if accepted == "Yes" else stage
        rows.append([f"A26-{i:04d}", applied, applied.strftime("%Y-%m"), s, src, rec_names[rec], final, offer_made,
                     accepted, hire_date, tth, SOURCE_COST[src]])
    rows.sort(key=lambda r: (r[1], r[0]))
    for i, r in enumerate(rows, 1):
        r[0] = f"A26-{i:04d}"
    headers = ["Application_ID", "Applied_Date", "Applied_Month", "Sector", "Source", "Recruiter", "Furthest_Stage",
               "Offer_Made", "Offer_Accepted", "Hire_Date", "Time_to_Hire_Days", "Sourcing_Cost"]
    hires = [r for r in rows if r[6] == "Hired"]
    offers = [r for r in rows if r[7] == "Yes"]
    decided = [r for r in offers if r[8] in ("Yes", "No")]
    by_month = {}
    for r in hires:
        k = r[9].strftime("%Y-%m")
        by_month[k] = by_month.get(k, 0) + 1
    by_sector = {}
    for r in hires:
        by_sector[r[3]] = by_sector.get(r[3], 0) + 1
    by_rec = {}
    for r in hires:
        by_rec[r[5]] = by_rec.get(r[5], 0) + 1
    tech_rows = [r for r in rows if r[3] == "Technology"]
    tech_hires = [r for r in tech_rows if r[6] == "Hired"]
    exp = dict(apps=len(rows), hires=len(hires), offers=len(offers), decided=len(decided),
               accepted=sum(1 for r in offers if r[8] == "Yes"), pending=sum(1 for r in offers if r[8] == "Pending"),
               oar=round(sum(1 for r in offers if r[8] == "Yes") / len(decided) * 100, 1),
               avg_tth=round(statistics.mean(r[10] for r in hires), 1),
               total_cost=sum(r[11] for r in rows), cph=round(sum(r[11] for r in rows) / len(hires), 2),
               by_month=dict(sorted(by_month.items())), top_sector=max(by_sector, key=by_sector.get),
               by_sector=by_sector, top_rec=max(by_rec, key=by_rec.get), top_rec_n=max(by_rec.values()),
               tech_apps=len(tech_rows), tech_hires=len(tech_hires),
               tech_avg_tth=round(statistics.mean(r[10] for r in tech_hires), 1) if tech_hires else 0)
    targets = [[m, t] for m, t in zip(["2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"],
                                      [0, 8, 9, 9, 9, 9])]
    return headers, rows, targets, exp


def build_lab5(folder):
    headers, rows, targets, _ = _lab5()
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    _write_rows(ws, headers, rows, "tblApps", {"Sourcing_Cost": '"S$"#,##0'}, {"Sector": 18, "Source": 19, "Recruiter": 15})
    ws2 = wb.create_sheet("Targets")
    _write_rows(ws2, ["Hire_Month", "Hire_Target"], targets, "tblTargets")
    wb.create_sheet("Dashboard")
    _about_sheet(wb, "Recruitment KPI Dashboard Data (Jan–Jun 2026)",
                 "Purpose: one row per application in H1 2026 with the furthest stage reached and the offer outcome. "
                 "Use it in Lab 5 to build a recruitment KPI dashboard (KPI cards, PivotTables, charts and slicers) on "
                 "the Dashboard sheet with Copilot.",
                 [("Application_ID", "Text", "Unique key."), ("Applied_Date / Applied_Month", "Date / Text", "Application date and yyyy-mm label."),
                  ("Sector / Source / Recruiter", "Text", "Dashboard dimensions."),
                  ("Furthest_Stage", "Text", "Applied, Screened, Interviewed, Offered or Hired."),
                  ("Offer_Made", "Yes/No", "Whether an offer was extended."),
                  ("Offer_Accepted", "Yes/No/Pending", "Blank when no offer was made. Pending = decision after 2026-06-30."),
                  ("Hire_Date / Time_to_Hire_Days", "Date / Days", "Filled only for hires."),
                  ("Sourcing_Cost", "S$", "Channel cost attributed to this application."),
                  ("KPI definitions", "Rule", "Offer acceptance rate = accepted ÷ (accepted + declined); Cost per hire = total "
                   "Sourcing_Cost ÷ hires; Time to hire = average Time_to_Hire_Days of hires; hires are counted by Hire_Date month.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Recruitment_Dashboard_Data.xlsx")]


# =========================================================================== LAB 6
DEPTS = ["Recruitment", "Sales & Client Services", "Payroll Services", "Finance", "Human Resources", "IT", "Operations"]


def _active_on(r, d):
    return r[5] <= d and (r[6] is None or r[6] > d)


def _lab6():
    rng = random.Random(3406)
    locs = ["Raffles Place HQ", "Jurong East Hub", "Tampines Hub", "Woodlands Hub"]
    dept_w = [30, 18, 14, 8, 7, 8, 15]
    level_w = {"Associate": 45, "Senior Associate": 28, "Manager": 18, "Senior Manager": 7, "Director": 2}
    level_sal = {"Associate": 3600, "Senior Associate": 5000, "Manager": 7400, "Senior Manager": 10200, "Director": 15500}
    rows = []
    for i in range(1, 421):
        d = rng.choices(DEPTS, dept_w)[0]
        loc = rng.choices(locs, [50, 18, 18, 14])[0]
        lvl = rng.choices(list(level_w), list(level_w.values()))[0]
        et = rng.choices(["Full-time", "Part-time", "Contract"], [80, 8, 12])[0]
        hire = _rand_date(rng, date(2019, 1, 7), date(2026, 6, 1))
        hazard = 0.26 if d in ("Recruitment", "Sales & Client Services") else 0.11
        exit_d, etype, status = None, "", "Active"
        cand = hire + timedelta(days=max(90, int(rng.expovariate(hazard) * 365)))
        if cand <= REPORT_DATE:
            exit_d = cand
            etype = rng.choices(["Voluntary", "Involuntary"], [82, 18])[0]
            status = "Exited"
        sal = int(round(level_sal[lvl] * rng.uniform(0.88, 1.15), -1))
        eng = round(min(5, max(1, rng.gauss(3.7 if status == "Active" else 3.0, 0.6))), 1)
        trn = max(0, int(rng.gauss(18, 8)))
        rows.append([f"EMP-{i:04d}", d, loc, lvl, et, hire, exit_d, etype, status, sal, eng, trn])
    rows.sort(key=lambda r: r[5])
    for i, r in enumerate(rows, 1):
        r[0] = f"EMP-{i:04d}"
    headers = ["Employee_ID", "Department", "Location", "Job_Level", "Employment_Type", "Hire_Date", "Exit_Date",
               "Exit_Type", "Status", "Monthly_Salary", "Engagement_Score", "Training_Hours_YTD"]
    p0, p1 = date(2025, 7, 1), REPORT_DATE
    active = [r for r in rows if r[8] == "Active"]
    start_hc = sum(1 for r in rows if _active_on(r, p0))
    end_hc = len(active)
    exits = [r for r in rows if r[6] and p0 <= r[6] <= p1]
    vol = [r for r in exits if r[7] == "Voluntary"]
    avg_hc = (start_hc + end_hc) / 2
    hires_12m = [r for r in rows if p0 <= r[5] <= p1]
    by_dept = {}
    for d in DEPTS:
        s0 = sum(1 for r in rows if r[1] == d and _active_on(r, p0))
        s1 = sum(1 for r in active if r[1] == d)
        ex = sum(1 for r in exits if r[1] == d)
        by_dept[d] = dict(hc=s1, exits=ex, rate=round(ex / ((s0 + s1) / 2) * 100, 1) if (s0 + s1) else 0)
    tenure = statistics.mean((REPORT_DATE - r[5]).days / 365.25 for r in active)
    new_joiners = []
    rng2 = random.Random(34061)
    for j in range(1, 9):
        d = rng2.choice(DEPTS[:3] + ["IT"])
        new_joiners.append([f"EMP-{len(rows) + j:04d}", d, rng2.choice(locs), "Associate", "Full-time",
                            date(2026, 7, rng2.choice([1, 6, 13, 20])), None, "", "Active",
                            int(round(3600 * rng2.uniform(0.9, 1.1), -1)), None, 0])
    exp = dict(rows=len(rows), active=end_hc, start_hc=start_hc, exits_12m=len(exits), vol_exits=len(vol),
               attrition=round(len(exits) / avg_hc * 100, 1), vol_attrition=round(len(vol) / avg_hc * 100, 1),
               hires_12m=len(hires_12m), by_dept=by_dept,
               top_attr_dept=max(by_dept, key=lambda k: by_dept[k]["rate"]), avg_tenure=round(tenure, 1),
               largest_dept=max(by_dept, key=lambda k: by_dept[k]["hc"]), new_joiners=len(new_joiners),
               avg_eng_active=round(statistics.mean(r[10] for r in active), 2),
               nj_by_dept={d: sum(1 for r in new_joiners if r[1] == d) for d in DEPTS if any(r[1] == d for r in new_joiners)},
               active_after=end_hc + len(new_joiners), rows_after=len(rows) + len(new_joiners))
    return headers, rows, new_joiners, exp


def build_lab6(folder):
    headers, rows, new_joiners, _ = _lab6()
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    _write_rows(ws, headers, rows, "tblEmployees", {"Monthly_Salary": '"S$"#,##0', "Engagement_Score": "0.0"},
                {"Department": 22, "Location": 18, "Job_Level": 16})
    ws2 = wb.create_sheet("New_Joiners_Jul2026")
    _write_rows(ws2, headers, new_joiners, None, {"Monthly_Salary": '"S$"#,##0'}, {"Department": 22, "Location": 18})
    ws2["N1"] = "Do NOT use until Lab 6 Step 7 — these rows are appended to tblEmployees to test dashboard refresh."
    ws2["N1"].font = Font(bold=True, color="C00000")
    wb.create_sheet("HR_Dashboard")
    _about_sheet(wb, "Internal Workforce Snapshot (as at 30 Jun 2026)",
                 "Purpose: one row per BrightPath employee hired since January 2019 (Status Active or Exited). Use it in "
                 "Lab 6 to build an interactive HR workforce and attrition dashboard. Salary and engagement data are "
                 "sensitive: share only aggregated views and never single out an individual.",
                 [("Employee_ID", "Text", "Pseudonymised employee key (no names are stored)."),
                  ("Department / Location / Job_Level / Employment_Type", "Text", "Dashboard dimensions."),
                  ("Hire_Date / Exit_Date", "Date", "Exit_Date blank for active employees."),
                  ("Exit_Type", "Text", "Voluntary or Involuntary; blank if active."), ("Status", "Text", "Active or Exited."),
                  ("Monthly_Salary", "S$", "Monthly base salary."), ("Engagement_Score", "1–5", "Latest pulse-survey score."),
                  ("Training_Hours_YTD", "Hours", "Training hours this calendar year."),
                  ("KPI definitions", "Rule", "12-month period = 2025-07-01 to 2026-06-30. Average headcount = (active on "
                   "2025-07-01 + active on 2026-06-30) ÷ 2. Attrition rate = exits in period ÷ average headcount.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Workforce_Headcount.xlsx")]


# =========================================================================== LAB 7
def _lab7():
    import math
    rng = random.Random(3407)
    rows = []
    for i in range(1, 401):
        d = rng.choices(DEPTS, [30, 18, 14, 8, 7, 8, 15])[0]
        lvl = rng.choices(["Associate", "Senior Associate", "Manager", "Senior Manager"], [48, 28, 18, 6])[0]
        tenure = round(max(0.2, rng.expovariate(1 / 3.6)), 1)
        mode = rng.choices(["Hybrid", "On-site", "Remote"], [55, 35, 10])[0]
        ot = max(0, int(rng.gauss(14 if d in ("Recruitment", "Sales & Client Services") else 8, 7)))
        eng = round(min(5, max(1, rng.gauss(3.6 - 0.03 * max(0, ot - 10), 0.65))), 1)
        promo = round(min(tenure, max(0, rng.gauss(2.2, 1.4))), 1)
        compa = round(rng.gauss(0.98, 0.09), 2)
        trn = max(0, int(rng.gauss(20, 9)))
        perf = rng.choices([1, 2, 3, 4, 5], [3, 10, 47, 30, 10])[0]
        mgr = f"MGR-{rng.randint(1, 24):02d}"
        z = (-2.1 + 0.07 * (ot - 10) - 0.95 * (eng - 3.5) + 0.32 * max(0, promo - 2.5)
             - 3.0 * (compa - 1.0) - 0.015 * (trn - 20) - 0.12 * min(tenure, 6) + 0.35 * (mode == "On-site"))
        left = "Yes" if rng.random() < 1 / (1 + math.exp(-z)) else "No"
        rows.append([f"EMP-{7000 + i}", d, lvl, tenure, mode, ot, eng, promo, compa, trn, perf, mgr, left])
    headers = ["Employee_ID", "Department", "Job_Level", "Tenure_Years", "Work_Mode", "Overtime_Hours_Month",
               "Engagement_Score", "Years_Since_Promotion", "Compa_Ratio", "Training_Hours", "Performance_Rating",
               "Manager_ID", "Left_Company"]

    def rate(rs):
        return round(sum(1 for r in rs if r[12] == "Yes") / len(rs) * 100, 1) if rs else 0
    by_dept = {d: rate([r for r in rows if r[1] == d]) for d in DEPTS}
    eng_lo, eng_hi = [r for r in rows if r[6] < 3], [r for r in rows if r[6] >= 3]
    ot_hi, ot_lo = [r for r in rows if r[5] > 20], [r for r in rows if r[5] <= 20]
    pr_hi, pr_lo = [r for r in rows if r[7] >= 3], [r for r in rows if r[7] < 3]
    cr_lo, cr_hi = [r for r in rows if r[8] < 0.9], [r for r in rows if r[8] >= 0.9]
    by_mode = {m: rate([r for r in rows if r[4] == m]) for m in ("Hybrid", "On-site", "Remote")}
    left = [r for r in rows if r[12] == "Yes"]
    stayed = [r for r in rows if r[12] == "No"]
    exp = dict(rows=len(rows), left=len(left), rate=rate(rows), by_dept=by_dept,
               dept_n={d: sum(1 for r in rows if r[1] == d) for d in DEPTS},
               top_dept=max(by_dept, key=by_dept.get), low_dept=min(by_dept, key=by_dept.get),
               eng_lo=(len(eng_lo), rate(eng_lo)), eng_hi=(len(eng_hi), rate(eng_hi)),
               ot_hi=(len(ot_hi), rate(ot_hi)), ot_lo=(len(ot_lo), rate(ot_lo)),
               pr_hi=(len(pr_hi), rate(pr_hi)), pr_lo=(len(pr_lo), rate(pr_lo)),
               cr_lo=(len(cr_lo), rate(cr_lo)), cr_hi=(len(cr_hi), rate(cr_hi)), by_mode=by_mode,
               eng_left=round(statistics.mean(r[6] for r in left), 2), eng_stayed=round(statistics.mean(r[6] for r in stayed), 2),
               ot_left=round(statistics.mean(r[5] for r in left), 1), ot_stayed=round(statistics.mean(r[5] for r in stayed), 1),
               corr_eng_ot=round(statistics.correlation([r[5] for r in rows], [r[6] for r in rows]), 2))
    return headers, rows, exp


def build_lab7(folder):
    headers, rows, _ = _lab7()
    wb = Workbook()
    ws = wb.active
    ws.title = "Attrition_Data"
    _write_rows(ws, headers, rows, "tblAttrition", {"Engagement_Score": "0.0", "Compa_Ratio": "0.00",
                                                    "Tenure_Years": "0.0", "Years_Since_Promotion": "0.0"},
                {"Department": 22, "Job_Level": 16})
    ws2 = wb.create_sheet("Insight_Log")
    _write_rows(ws2, ["Insight_No", "Observation (what the data shows)", "Evidence (numbers + how verified)",
                      "Label (OBSERVED / HYPOTHESIS)", "Possible explanation to test", "Limitation"], [], None,
                widths={"Observation (what the data shows)": 45, "Evidence (numbers + how verified)": 45,
                        "Possible explanation to test": 40, "Limitation": 34})
    for c in range(1, 7):
        ws2.cell(row=1, column=c).fill = HEADER_FILL
        ws2.cell(row=1, column=c).font = HEADER_FONT
    _about_sheet(wb, "Employee Attrition Study (Jul 2025 – Jun 2026)",
                 "Purpose: one row per BrightPath employee who was active on 2025-07-01, with work-pattern and "
                 "engagement attributes and whether they left within the following 12 months. Use it in Lab 7 to "
                 "surface insights with Copilot and Python-based advanced analysis. Treat every pattern as a team-level "
                 "signal: never use this data to make a decision about an individual.",
                 [("Employee_ID / Manager_ID", "Text", "Pseudonymised keys."),
                  ("Department / Job_Level / Work_Mode", "Text", "Grouping attributes."),
                  ("Tenure_Years", "Years", "Tenure on 2025-07-01."),
                  ("Overtime_Hours_Month", "Hours", "Average monthly overtime over the prior 6 months."),
                  ("Engagement_Score", "1–5", "Latest engagement pulse score."),
                  ("Years_Since_Promotion", "Years", "Time since last promotion or hire."),
                  ("Compa_Ratio", "Ratio", "Salary ÷ midpoint of the salary band (1.00 = at midpoint)."),
                  ("Training_Hours", "Hours", "Training hours in the prior 12 months."),
                  ("Performance_Rating", "1–5", "Latest rating."),
                  ("Left_Company", "Yes/No", "Whether the employee left between 2025-07-01 and 2026-06-30."),
                  ("Insight_Log sheet", "Template", "Record each observation, its evidence and its label.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Attrition_Analysis.xlsx")]


# =========================================================================== LAB 8
QUESTIONS = ["Q1_Career_Growth", "Q2_Manager_Support", "Q3_Workload_Balance", "Q4_Pay_Fairness",
             "Q5_Recognition", "Q6_Tools_and_Systems"]
COMMENTS = {
    "Workload": [("Desk targets keep rising and we are always short of hands during peak hiring.", "Negative"),
                 ("Too many late nights to close client requisitions before month end.", "Negative"),
                 ("Workload is heavy but manageable when the team shares candidates.", "Mixed")],
    "Career Growth": [("I don't see a clear path from Senior Associate to Manager.", "Negative"),
                      ("Would like structured mentoring and rotation across desks.", "Mixed"),
                      ("Promotions feel slow compared with other agencies.", "Negative")],
    "Manager Support": [("My manager gives quick feedback and backs us with difficult clients.", "Positive"),
                        ("Weekly one-on-ones are often cancelled.", "Negative"),
                        ("Team lead is supportive and approachable.", "Positive")],
    "Pay & Benefits": [("Commission scheme is hard to understand and changes too often.", "Negative"),
                       ("Base pay is below market for experienced recruiters.", "Negative"),
                       ("Benefits are fair and the medical coverage is good.", "Positive")],
    "Recognition": [("Wins are celebrated in the town hall, which is motivating.", "Positive"),
                    ("Placements we close are rarely recognised beyond the top biller.", "Negative")],
    "Tools & Systems": [("The applicant tracking system is slow and duplicates candidate records.", "Negative"),
                        ("New Copilot tools save time on job descriptions.", "Positive"),
                        ("Too many manual spreadsheets for payroll reconciliation.", "Negative")],
    "Flexibility": [("Hybrid work arrangement is a big reason I stay.", "Positive"),
                    ("Would appreciate more flexibility on in-office days.", "Mixed")],
}


def _lab8():
    rng = random.Random(3408)
    rows = []
    theme_w = {"Workload": 24, "Career Growth": 19, "Manager Support": 13, "Pay & Benefits": 16,
               "Recognition": 9, "Tools & Systems": 12, "Flexibility": 7}
    truth = []
    for i in range(1, 161):
        d = rng.choices(DEPTS, [30, 18, 14, 8, 7, 8, 15])[0]
        tb = rng.choices(["< 1 year", "1–3 years", "3–5 years", "5+ years"], [18, 36, 26, 20])[0]
        wm = rng.choices(["Hybrid", "On-site", "Remote"], [55, 35, 10])[0]
        dept_adj = -0.35 if d in ("Recruitment", "Sales & Client Services") else 0.1
        base = {"Q1_Career_Growth": 3.0, "Q2_Manager_Support": 3.8, "Q3_Workload_Balance": 2.8,
                "Q4_Pay_Fairness": 3.1, "Q5_Recognition": 3.3, "Q6_Tools_and_Systems": 3.2}
        scores = [int(min(5, max(1, round(rng.gauss(base[q] + dept_adj * (q in ("Q3_Workload_Balance", "Q1_Career_Growth")), 0.85)))))
                  for q in QUESTIONS]
        mean_s = sum(scores) / len(scores)
        enps = int(min(10, max(0, round(rng.gauss(2.4 * mean_s - 0.6, 1.4)))))
        comment, theme = "", ""
        if rng.random() < 0.85:
            theme = rng.choices(list(theme_w), list(theme_w.values()))[0]
            comment = rng.choice(COMMENTS[theme])[0]
        truth.append(theme)
        rows.append([f"SR-{i:03d}", d, tb, wm] + scores + [enps, comment])
    headers = ["Response_ID", "Department", "Tenure_Band", "Work_Mode"] + QUESTIONS + ["eNPS_Score", "Comment"]

    def enps(rs):
        if not rs:
            return 0
        p = sum(1 for r in rs if r[10] >= 9)
        dd = sum(1 for r in rs if r[10] <= 6)
        return round((p - dd) / len(rs) * 100, 1)
    q_avg = {q: round(statistics.mean(r[4 + i] for r in rows), 2) for i, q in enumerate(QUESTIONS)}
    by_dept = {d: enps([r for r in rows if r[1] == d]) for d in DEPTS}
    counts = {}
    for t in truth:
        if t:
            counts[t] = counts.get(t, 0) + 1
    exp = dict(rows=len(rows), enps=enps(rows), promoters=sum(1 for r in rows if r[10] >= 9),
               passives=sum(1 for r in rows if 7 <= r[10] <= 8), detractors=sum(1 for r in rows if r[10] <= 6),
               q_avg=q_avg, lowest_q=min(q_avg, key=q_avg.get), highest_q=max(q_avg, key=q_avg.get),
               by_dept=by_dept, dept_n={d: sum(1 for r in rows if r[1] == d) for d in DEPTS},
               low_dept=min(by_dept, key=by_dept.get), comments=sum(1 for r in rows if r[11]),
               theme_counts=dict(sorted(counts.items(), key=lambda kv: -kv[1])),
               top_theme=max(counts, key=counts.get))
    return headers, rows, exp


def build_lab8(folder):
    headers, rows, _ = _lab8()
    wb = Workbook()
    ws = wb.active
    ws.title = "Survey_Responses"
    _write_rows(ws, headers, rows, "tblSurvey", None, {"Department": 22, "Tenure_Band": 12, "Comment": 70})
    ws2 = wb.create_sheet("Recommendations")
    _write_rows(ws2, ["Rec_No", "Finding (OBSERVED)", "Evidence (cells / numbers)", "Recommendation (action)",
                      "Owner", "Priority (High / Medium / Low)", "Success KPI and target", "Review date"], [], None,
                widths={"Finding (OBSERVED)": 40, "Evidence (cells / numbers)": 36, "Recommendation (action)": 44,
                        "Success KPI and target": 30})
    for c in range(1, 9):
        ws2.cell(row=1, column=c).fill = HEADER_FILL
        ws2.cell(row=1, column=c).font = HEADER_FONT
    wb.create_sheet("Executive_Summary")
    _about_sheet(wb, "Employee Engagement Pulse Survey (June 2026)",
                 "Purpose: one row per anonymous survey response from BrightPath staff. Use it in Lab 8 to quantify scores, "
                 "classify comment themes and sentiment, and turn verified findings into recommendations. Responses are "
                 "anonymous: do not try to identify respondents, and do not report any group with fewer than 5 responses.",
                 [("Response_ID", "Text", "Anonymous response key."),
                  ("Department / Tenure_Band / Work_Mode", "Text", "Grouping attributes."),
                  ("Q1 … Q6", "1–5", "Agreement score (1 = strongly disagree, 5 = strongly agree) for career growth, manager "
                   "support, workload balance, pay fairness, recognition, and tools and systems."),
                  ("eNPS_Score", "0–10", "'How likely are you to recommend BrightPath as a place to work?'"),
                  ("Comment", "Free text", "Optional comment; blank when the respondent left none."),
                  ("eNPS rule", "Rule", "eNPS = % promoters (9–10) − % detractors (0–6), ranging from −100 to +100."),
                  ("Recommendations / Executive_Summary", "Templates", "Output sheets for Steps 6–8.")])
    _prompt_log_sheet(wb)
    return [_save(wb, folder, "BrightPath_Engagement_Survey.xlsx")]


# =========================================================================== registry
_BUILDERS = {1: build_lab1, 2: build_lab2, 3: build_lab3, 4: build_lab4,
             5: build_lab5, 6: build_lab6, 7: build_lab7, 8: build_lab8}

EXPECTED = {1: _lab1()[2], 2: _lab2()[2], 3: _lab3()[3], 4: _lab4()[2],
            5: _lab5()[3], 6: _lab6()[3], 7: _lab7()[2], 8: _lab8()[2]}


def build_lab_data(num, folder):
    return _BUILDERS[num](folder) if num in _BUILDERS else []


if __name__ == "__main__":
    import json
    print(json.dumps(EXPECTED, indent=1, default=str))
